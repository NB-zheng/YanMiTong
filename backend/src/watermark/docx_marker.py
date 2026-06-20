"""Word 指纹嵌入：python-docx 写文档属性。"""
from __future__ import annotations

import io


def embed(docx_bytes: bytes, fingerprint: str) -> bytes:
    try:
        from docx import Document
    except ImportError:
        # 降级：docx 本质是 zip，简单追加 trailer 注释（非标准但取证可识别）
        return docx_bytes + f"\nYMT-FP:{fingerprint}\n".encode("utf-8")

    bio = io.BytesIO(docx_bytes)
    doc = Document(bio)
    cp = doc.core_properties
    cp.keywords = (cp.keywords or "") + f" YMTFingerprint({fingerprint})"
    cp.comments = (cp.comments or "") + f" YMT-FP:{fingerprint}"

    # 隐藏段落（字号极小、白色）
    p = doc.add_paragraph()
    run = p.add_run(f"YMT-FP:{fingerprint}")
    run.font.size = None  # 取默认；纯文本指纹方便取证 grep
    try:
        from docx.shared import RGBColor, Pt
        run.font.size = Pt(1)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    except Exception:
        pass

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
